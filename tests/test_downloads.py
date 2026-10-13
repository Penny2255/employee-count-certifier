# © 2026 Matthew Beltrano. All rights reserved. Proprietary and confidential. Not open source. See LICENSE.
"""Downloads: the 5-slide PowerPoint deck (light/dark), the Word document, their export routes, the links the
Claude connector returns, and the HQ address in the record."""

import io
import json
import os
from pathlib import Path

os.environ.setdefault("CERTIFIER_NO_APP", "1")

from docx import Document  # noqa: E402
from pptx import Presentation  # noqa: E402

from certifier import office  # noqa: E402
from certifier.record import record_from_lookup  # noqa: E402
from certifier.service import LookupResult  # noqa: E402

from .test_web import INIT, SF, _mcp, _oauth_token, auth, client, env, result_of, stream  # noqa: F401,E402

SAMPLE = Path(__file__).resolve().parents[1] / "samples" / "salesforce.json"
LINKEDIN = "https://www.linkedin.com/in/matthewbeltrano/"


def record():
    return json.loads(record_from_lookup(LookupResult.model_validate_json(SAMPLE.read_text(encoding="utf-8"))).model_dump_json())


def test_deck_has_five_branded_slides_with_company_site_and_credit_links_in_both_themes():
    site = "https://employee-count-certifier.onrender.com"
    for theme in ("light", "dark"):
        prs = Presentation(io.BytesIO(office.deck(record(), theme, site_url=site)))
        assert len(prs.slides) == 5
        titles = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]
        for word, slide in zip(("The answer", "Trend and forecast", "Why you can trust this", "Sources", "Thank you."), titles):
            assert word in slide
        assert all("Salesforce, Inc." in slide for slide in titles)  # the company on every slide
        assert "83,334" in titles[0] and "Designed and created by" in titles[0]
        assert "elevate human needs and outcomes" in titles[4]
        links = [r.hyperlink.address for s in prs.slides for sh in s.shapes if sh.has_text_frame
                 for p in sh.text_frame.paragraphs for r in p.runs if r.hyperlink.address]
        links += [sh.click_action.hyperlink.address for s in prs.slides for sh in s.shapes
                  if sh.click_action.hyperlink.address]  # clickable buttons and the LinkedIn card
        links += [r.hyperlink.address for s in prs.slides for sh in s.shapes if sh.has_table for row in sh.table.rows
                  for c in row.cells for p in c.text_frame.paragraphs for r in p.runs if r.hyperlink.address]  # table cells
        assert links.count(LINKEDIN) >= 5 and links.count(site) >= 5 and any("sec.gov" in u for u in links)
        assert any(sh.has_chart for sh in prs.slides[1].shapes)  # a native, editable chart
        card = [sh for sh in prs.slides[4].shapes if sh.shape_type == 13 and sh.click_action.hyperlink.address == LINKEDIN]
        assert card  # the LinkedIn card links to the profile


def test_word_document_has_the_detail_and_the_credit_link():
    doc = Document(io.BytesIO(office.document(record())))
    text = "\n".join(p.text for p in doc.paragraphs)
    for heading in ("Trend and forecast", "Findings", "Key source", "Sources", "Events and LinkedIn", "How it was certified"):
        assert heading in text
    assert "Designed and created by Matthew Beltrano & Claude Code" in text
    assert any(rel.target_ref == LINKEDIN for rel in doc.part.rels.values())


def test_export_routes_for_slides_and_document(client):
    qid = result_of(stream(client, SF)[1])["query_id"]
    deck = client.get(f"/api/lookups/{qid}/export?format=pptx&theme=dark")
    assert deck.status_code == 200 and deck.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.presentationml")
    assert 'filename="salesforce-inc-certification-dark.pptx"' in deck.headers["content-disposition"]
    assert len(Presentation(io.BytesIO(deck.content)).slides) == 5
    doc = client.get(f"/api/lookups/{qid}/export?format=docx")
    assert doc.status_code == 200 and 'certification.docx"' in doc.headers["content-disposition"]
    assert client.get("/api/lookups/sample-salesforce/export?format=pptx").status_code == 200  # samples too


def test_connector_results_carry_download_links(client, env, monkeypatch):
    from certifier import server as mcp_server
    monkeypatch.setattr(mcp_server, "lookup", __import__("certifier.web", fromlist=["lookup"]).lookup)
    token = _oauth_token(client, mode="guest")["access_token"]
    _mcp(client, token, "initialize", INIT)
    _, out = _mcp(client, token, "tools/call", {"name": "certify_employee_count",
                                                "arguments": {"name": "Salesforce, Inc.", "website": "https://www.salesforce.com"}}, 2)
    sc = out["result"]["structuredContent"]
    links = sc.get("result", sc)["entity"]["downloads"]
    assert set(links) == {"json", "summary_csv", "sources_csv", "pptx_light", "pptx_dark", "docx"}
    url = links["pptx_dark"]
    assert url.startswith("http://localhost/api/lookups/") and url.endswith("format=pptx&theme=dark")
    assert client.get(url.removeprefix("http://localhost")).status_code == 200


def test_record_carries_the_registered_hq_address():
    r = record()
    assert r["entity"]["address"].startswith("Salesforce Tower, 415 Mission Street 3rd Fl")
    assert r["entity"]["address_source"] == "SEC EDGAR (registered business address)"
